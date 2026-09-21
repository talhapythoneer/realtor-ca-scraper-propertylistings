"""Reading/writing the master, fresh and excluded-listings output files.

One master + one fresh CSV/XLSX pair per region:
  - master: every listing ever captured for that region, deduplicated by MLS
    number. Grows across runs and is the source of truth for "already sent".
  - fresh:  only the new, non-excluded listings found during *this* run.
    This is what gets handed to the print shop for that week's mailer.

Excluded listings (land/investor/teardown matches) are appended to a single
cumulative log across all regions, so the client can review them and use the
list later to fine-tune an AI-based filter, per the discussion with the client.
"""
import csv
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook, load_workbook

from .address import parse_address
from .models import OUTPUT_COLUMNS, Listing
from .numeric import clean_bedrooms, clean_leading_number, clean_price

logger = logging.getLogger("realtor_scraper")

# Fields recomputed from full_address on every write, so historical rows scraped
# before this parsing existed get backfilled automatically instead of staying blank.
_ADDRESS_DERIVED_FIELDS = ("unit", "street_address", "province", "postal_code")

# Numeric display fields ("$698,000", "3 + 2", "1200+") converted to real numbers so
# Excel/CSV consumers can sort and filter them, instead of treating them as text.
_NUMERIC_FIELD_CLEANERS = {
    "price": clean_price,
    "bedrooms": clean_bedrooms,
    "bathrooms": clean_leading_number,
    "square_footage": clean_leading_number,
    "storeys": clean_leading_number,
}

# What the print shop actually needs, per the client's spec: address (split out),
# price, building type, date listed - not the full raw scrape.
PRINT_SHOP_COLUMNS = [
    "unit",
    "street_address",
    "city",
    "province",
    "postal_code",
    "price",
    "building_type",
    "estimated_listed_date",
]
PRINT_SHOP_HEADERS = [
    "Unit",
    "Street Address",
    "City",
    "Province",
    "Postal Code",
    "Price",
    "Building Type",
    "Date Listed",
]


def normalize_row(row: dict) -> dict:
    """Recompute the address-split and numeric columns for one output row.

    Applied to every row at write time (not just newly-scraped ones), so the
    master file self-heals: rows written by an older version of the scraper
    (blank unit/province, "$698,000"-style price, etc.) get cleaned up the
    next time they're written out, with no separate migration step needed.
    """
    row = dict(row)
    parsed = parse_address(row.get("full_address", ""))
    for field in _ADDRESS_DERIVED_FIELDS:
        row[field] = parsed[field]
    for field, cleaner in _NUMERIC_FIELD_CLEANERS.items():
        value = cleaner(row.get(field, ""))
        row[field] = value if value is not None else ""
    return row


def _safe_region_filename(region: str) -> str:
    return "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in region).strip().replace(" ", "_")


def master_path(output_dir: Path, region: str, template: str, ext: str) -> Path:
    filename = template.format(region=_safe_region_filename(region), ext=ext)
    return output_dir / filename


def fresh_path(output_dir: Path, region: str, template: str, ext: str, timestamp: str) -> Path:
    filename = template.format(region=_safe_region_filename(region), ext=ext, timestamp=timestamp)
    return output_dir / filename


def _load_mls_numbers(csv_path: Path) -> set:
    if not csv_path.exists():
        return set()
    seen = set()
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            mls = (row.get("mls_number") or "").strip()
            if mls:
                seen.add(mls)
            elif (row.get("listing_id") or "").strip():
                seen.add("id:" + row["listing_id"].strip())
    return seen


def load_master_mls_numbers(csv_path: Path) -> set:
    """Return the set of mls_number values already present in a region's master CSV."""
    return _load_mls_numbers(csv_path)


def load_excluded_mls_numbers(excluded_log_path: Path) -> set:
    """Return the set of mls_number values already logged as excluded.

    Checked alongside the master file so a listing that was excluded once
    (land/investor/teardown keyword match) isn't re-fetched and re-checked
    on every subsequent run - it's already been judged.
    """
    return _load_mls_numbers(excluded_log_path)


def _write_csv(path: Path, rows: List[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _cell_value(col: str, value):
    """Give date columns a real date type so Excel sorts/filters them as dates.
    Numeric columns are already real int/float by the time rows reach here
    (normalize_row converts them) - openpyxl writes those as-is."""
    if value in (None, ""):
        return None
    if col == "estimated_listed_date" and isinstance(value, str):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return value
    return value


def _write_sheet(ws, columns: List[str], headers: List[str], rows: List[dict]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_cell_value(col, row.get(col, "")) for col in columns])
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, min(40, len(header) + 4))
    ws.freeze_panes = "A2"


def _write_xlsx(path: Path, rows: List[dict]) -> None:
    """Two tabs: "Raw Data" (everything scraped) and "Print Shop" (just what the
    mailer run needs - address, price, building type, date listed), per the client's
    request. (A plain .csv has no concept of tabs, so this split only applies to .xlsx -
    the .csv companion file stays the single raw table it's always been.)"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    raw_ws = wb.active
    raw_ws.title = "Raw Data"
    _write_sheet(raw_ws, OUTPUT_COLUMNS, OUTPUT_COLUMNS, rows)

    print_ws = wb.create_sheet("Print Shop")
    _write_sheet(print_ws, PRINT_SHOP_COLUMNS, PRINT_SHOP_HEADERS, rows)

    wb.save(path)


def read_master_rows(csv_path: Path) -> List[dict]:
    if not csv_path.exists():
        return []
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return [row for row in reader]


def write_region_outputs(
    output_dir: Path,
    region: str,
    new_listings: List[Listing],
    formats: List[str],
    master_template: str,
    fresh_template: str,
    timestamp: str,
) -> Dict[str, Path]:
    """Append new_listings to the region's master file(s) and write a fresh file
    containing just this run's new listings. Returns a dict of written paths."""
    written: Dict[str, Path] = {}
    new_rows = [normalize_row(listing.to_row()) for listing in new_listings]

    # The CSV master is always kept up to date, even if "csv" isn't in the
    # configured output formats - it's the dedup source of truth that
    # load_master_mls_numbers() reads on the next run.
    csv_master_path = master_path(output_dir, region, master_template, "csv")
    combined_rows = [normalize_row(r) for r in read_master_rows(csv_master_path)] + new_rows
    _write_csv(csv_master_path, combined_rows)
    if "csv" in formats:
        written["master_csv"] = csv_master_path

    if "xlsx" in formats:
        xlsx_master_path = master_path(output_dir, region, master_template, "xlsx")
        _write_xlsx(xlsx_master_path, combined_rows)
        written["master_xlsx"] = xlsx_master_path

    for ext in formats:
        f_path = fresh_path(output_dir, region, fresh_template, ext, timestamp)
        if ext == "csv":
            _write_csv(f_path, new_rows)
        elif ext == "xlsx":
            _write_xlsx(f_path, new_rows)
        else:
            continue
        written[f"fresh_{ext}"] = f_path

    return written


def append_excluded_log(path: Path, excluded_listings: List[Listing]) -> None:
    """Append excluded listings to a cumulative CSV log, skipping MLS numbers already logged.

    Rewrites the whole file (existing rows + new ones) rather than a raw
    file-append, same as the master CSV - this keeps the header in sync with
    OUTPUT_COLUMNS even as it evolves (e.g. the "unit" column added later),
    and backfills address/numeric fields on existing rows via normalize_row().
    A blind append under a stale header would silently misalign columns.
    """
    if not excluded_listings:
        return

    existing_rows = [normalize_row(r) for r in read_master_rows(path)]
    existing_mls = {(r.get("mls_number") or "").strip() for r in existing_rows} - {""}

    new_rows = [
        normalize_row(listing.to_row())
        for listing in excluded_listings
        if not (listing.mls_number and listing.mls_number in existing_mls)
    ]
    if not new_rows:
        return

    _write_csv(path, existing_rows + new_rows)
